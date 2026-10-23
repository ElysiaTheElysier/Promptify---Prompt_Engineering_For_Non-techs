from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "document-drafts"
SHA = "bf1d194d67524683776606095fe13fa338252c12"
BASE = f"https://github.com/ElysiaTheElysier/Promptify---Prompt_Engineering_For_Non-techs/blob/{SHA}/"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); tc_pr.append(el)

def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    bd = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'4'); tag.set(qn('w:color'),'D9D9D9')
        bd.append(tag)
    tc_pr.append(bd)

def make_doc(title, subtitle):
    d = Document()
    sec = d.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.58)
    sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
    styles = d.styles
    normal = styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10); normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1.06
    for name,size,before,after in [('Title',18,0,8),('Heading 1',12,10,5),('Heading 2',10.5,7,3)]:
        s=styles[name]; s.font.name='Aptos Display' if name=='Title' else 'Aptos'; s.font.size=Pt(size)
        s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
        s.paragraph_format.keep_with_next=True
        if name=='Title':
            ppr=s._element.get_or_add_pPr()
            for node in ppr.findall(qn('w:pBdr')): ppr.remove(node)
    styles['Subtitle'].font.color.rgb=RGBColor(0,0,0)
    styles['Subtitle'].font.italic=False
    if 'Small Note' not in styles:
        s=styles.add_style('Small Note',WD_STYLE_TYPE.PARAGRAPH); s.font.name='Aptos'; s.font.size=Pt(8.5)
        s.font.color.rgb=RGBColor(75,85,99); s.paragraph_format.space_after=Pt(5)
    d.add_paragraph(title,'Title')
    d.add_paragraph(subtitle,'Subtitle')
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Promptify  |  03/10/2026  |  ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    return d

def h(d,s,l=1): d.add_heading(s,l)
def p(d,s,style=None): d.add_paragraph(s,style)
def bullets(d,items):
    for x in items: d.add_paragraph(x,'List Bullet')
def nums(d,items):
    for i,x in enumerate(items,1):
        para=d.add_paragraph(f'{i}. {x}')
        para.paragraph_format.left_indent=Inches(.22)
        para.paragraph_format.first_line_indent=Inches(-.22)
def table(d,heads,rows,widths=None):
    t=d.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,v in enumerate(heads):
        c=t.rows[0].cells[i]; c.text=v; shade(c,'E9EEF3'); border(c)
        for run in c.paragraphs[0].runs: run.bold=True
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            border(cells[i]);
            if ri%2: shade(cells[i],'F8FAFC')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr()
        if not trpr.findall(qn('w:cantSplit')): trpr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                for run in para.runs: run.font.size=Pt(8.4)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    d.add_paragraph()
    return t

def source(d,entries):
    h(d,'Nguồn kiểm chứng')
    paths='; '.join(path for _,path in entries)
    p(d,f'GitHub: https://github.com/ElysiaTheElysier/Promptify---Prompt_Engineering_For_Non-techs  |  main {SHA}. File đối chiếu: {paths}. Chưa đối chiếu được commit deploy Vercel.','Small Note')

def brd():
    d=make_doc('Yêu cầu nghiệp vụ Promptify','BRD hiện trạng cho quản lý dự án và giảng viên chỉnh sửa')
    p(d,'Tài liệu này xác định phạm vi nghiệp vụ đang thể hiện trong mã Promptify main: người học thực hành prompt theo lớp, giảng viên quản lý nội dung và theo dõi tiến độ. Người quản lý dự án dùng tài liệu để thống nhất luồng nghiệp vụ và đánh dấu yêu cầu cần thay đổi. Các mục chưa có bằng chứng được để thành quyết định mở, không xem là chức năng hiện hữu.')
    table(d,['Thông tin','Giá trị'],[['Phiên bản nguồn',f'main tại commit {SHA[:12]}'],['Trạng thái','Bản nháp as-is, chờ thầy duyệt/chỉnh sửa'],['Deploy','https://promptify-wheat-seven.vercel.app/; chưa xác minh commit đang chạy']], [1.55,5.3])
    h(d,'1 Bối cảnh và mục tiêu nghiệp vụ')
    p(d,'Promptify phục vụ học viên không chuyên kỹ thuật thực hành viết prompt với tình huống công việc. Luồng sản phẩm cho phép đọc đề bài, viết và chạy prompt, nhận kết quả và đánh giá, cải thiện qua các lần thử. Giảng viên có giao diện để theo dõi lớp, học viên và biên tập chương trình học. Đây là mục tiêu suy từ giao diện và mã; mục tiêu định lượng, quy mô triển khai và chính sách đào tạo cần chủ dự án xác nhận.')
    h(d,'2 Bên liên quan và vai trò')
    table(d,['Vai trò','Nhu cầu trong luồng hiện tại','Ranh giới'],[
        ['Học viên learner','Vào lớp, học bài, thực hành prompt, xem kết quả và tiến độ','Cần phiên đăng nhập và quyền ghi danh lớp; lớp self_enroll có thể tự tham gia'],
        ['Giảng viên instructor','Quản lý khóa học/nội dung, danh sách học viên, quan sát hoạt động và tiến độ','Phạm vi dữ liệu phụ thuộc truy vấn và RLS Supabase'],
        ['Admin','Mã có role admin và đưa vào giao diện giảng viên','Chưa thấy màn hoặc bộ quyền admin riêng; cần quyết định nghiệp vụ']
    ],[1.25,3.1,2.5])
    h(d,'3 Phạm vi nghiệp vụ hiện có')
    table(d,['Nhóm năng lực','Đầu ra nghiệp vụ'],[
        ['Truy cập','Đăng nhập Google; phân vai; xác nhận ghi danh trước khi vào không gian học'],
        ['Tổ chức đào tạo','Course chứa module và lesson; class gắn course với client/department; learner tham gia class qua enrollment'],
        ['Học tập','Xem lộ trình, bài xuất bản, lý thuyết, bài tập, gợi ý và lời giải'],
        ['Thực hành','Soạn và chạy prompt, nhận output, đánh giá năm tiêu chí, so sánh lần thử'],
        ['Giám sát','Giảng viên xem lớp, học viên, hoạt động, attempt và progress'],
        ['Quản trị nội dung','Tạo/sửa/xuất bản khóa học; quản lý module, lesson, rubric, resource']
    ],[1.65,5.2])
    h(d,'4 Luồng nghiệp vụ từ đầu đến cuối')
    h(d,'4.1 Học viên',2)
    nums(d,[
        'Mở trang công khai và đăng nhập Google OAuth. Ứng dụng đồng bộ hồ sơ người dùng và lấy role từ public.users.role.',
        'Hệ thống lấy enrollment còn hiệu lực. Nếu chưa có, hiển thị chọn lớp; lớp self_enroll cho phép tự ghi danh.',
        'Học viên chọn lớp, xem trang chủ và lộ trình của course. Bài học được lấy từ nội dung published của course khi dữ liệu sẵn có.',
        'Trong bài học HybridView, học viên đọc lý thuyết và tình huống, soạn prompt, chạy AI thật hoặc mô phỏng nếu chủ động chọn.',
        'Ứng dụng trả output, chấm theo rubric và cho so sánh/cải thiện prompt. Lần thử và tiến độ được lưu sau khi ghi nhận thành công.',
        'Học viên xem thư viện prompt và lịch sử trong trình duyệt; màn lịch sử này có chức năng xuất Markdown.'
    ])
    h(d,'4.2 Giảng viên',2)
    nums(d,[
        'Đăng nhập; role instructor hoặc admin được đưa vào khu vực giảng viên.',
        'Xem dashboard, danh sách lớp, chi tiết lớp, số học viên, tiến độ theo bài và hoạt động.',
        'Thêm học viên bằng email vào lớp; sửa tên/trạng thái; đưa khỏi lớp bằng cách đổi enrollment sang removed.',
        'Tạo/sửa khóa học, đổi trạng thái xuất bản; quản lý chương, bài, rubric và tài nguyên.'
    ])
    h(d,'5 Quy tắc nghiệp vụ có bằng chứng')
    table(d,['Mã','Quy tắc hiện tại','Nguồn'],[
        ['BR01','Learner chưa ghi danh hợp lệ không được vào dashboard/bài học','App.tsx; dbService.ts'],
        ['BR02','Lớp self_enroll cho phép người học tự ghi danh','App.tsx; migration 012'],
        ['BR03','Người học chỉ thấy nội dung xuất bản và có quyền theo RLS; mã có fallback LABS_DATA khi tải curriculum lỗi','App.tsx; migration 006'],
        ['BR04','Đưa học viên khỏi lớp đổi trạng thái enrollment thành removed, không xóa user/learner','dbService.ts'],
        ['BR05','Bài học có thể có nhiều attempt; lesson_progress là trạng thái hiện tại theo learner, class, lesson','migration 016'],
        ['BR06','Admin hiện dùng cùng giao diện với instructor','App.tsx']
    ],[.7,4.65,1.5])
    h(d,'6 Mô hình thông tin nghiệp vụ')
    p(d,'Client (tổ chức) và Course (khóa học) được gắn vào Class (lớp). User có role learner sẽ có Learner profile; Learner tham gia Class qua Enrollment. Course gồm Course Module, mỗi module gồm Lesson; Lesson có rubric và resource. Một lần thực hành tạo Prompt Attempt và cập nhật Lesson Progress. Mô hình bảng và khóa chi tiết nằm ở PRD.')
    h(d,'7 Điểm phải quyết định trước khi nâng cấp')
    table(d,['Câu hỏi','Người quyết định / ghi chú'],[
        ['Mục tiêu chính của phiên bản tiếp theo và KPI có baseline?','[Thầy / PM điền]'],
        ['Admin có cần phân quyền độc lập với instructor?','[Thầy / PM điền]'],
        ['Lịch sử prompt cần xem xuyên thiết bị hay chỉ trên trình duyệt?','[Thầy / PM điền]'],
        ['Những course/lesson nào là baseline nghiệm thu?','[Thầy / PM điền]'],
        ['Commit nào đang chạy trên Vercel?','[Kỹ thuật xác nhận]']
    ],[4.4,2.45])
    h(d,'8 Ngoài phạm vi khẳng định của bản này')
    p(d,'Không có bằng chứng trong mã main cho cơ chế đồng bộ roster từ Google Sheets, Redis cache, KPI 45%/85%, hay quyền admin riêng. Repo có 22 tài liệu giáo trình nhưng không chứng minh 22 bài đều đã xuất bản cho mọi lớp. Các ý này nếu cần giữ phải được viết thành yêu cầu mới và được chủ dự án duyệt, không gắn nhãn “đã có”.')
    source(d,[('Luồng và vai trò','src/App.tsx'),('Mô hình lõi','supabase/schema.sql'),('Quản lý curriculum','supabase/migrations/004_course_content_management.sql'),('Truy cập theo enrollment','supabase/migrations/006_enforce_course_enrollment_access.sql'),('Tiến độ server','supabase/migrations/016_add_server_backed_lesson_progress.sql'),('Giảng viên','src/components/instructor/InstructorViewShell.tsx')])
    out=OUT/'BRD_Promptify_Quan_ly_du_an.docx'; d.save(out); return out

def prd():
    d=make_doc('Đặc tả sản phẩm Promptify','PRD hiện trạng cho quản lý dự án và giảng viên chỉnh sửa')
    p(d,'Tài liệu này mô tả đường chạy sản phẩm tại main để nhóm dự án dùng làm baseline: luồng màn hình, chức năng theo vai trò, dữ liệu, API và công nghệ. Nó không tự thêm tính năng tương lai. Các tiêu chí bên dưới là cách kiểm tra hành vi hiện có trong mã; việc nghiệm thu trên Vercel cần xác nhận commit deploy và dữ liệu môi trường.')
    table(d,['Thuộc tính','Giá trị'],[['Nguồn',f'GitHub main commit {SHA[:12]}'],['Phạm vi','Frontend, API, schema/migration và service hiện tại'],['Chưa xác nhận','Commit và cấu hình đang chạy trên Vercel']], [1.45,5.4])
    h(d,'1 Bản đồ màn hình và luồng người dùng')
    p(d,'Khách → Landing → Google OAuth → phân vai. Learner chưa có enrollment → chọn lớp. Learner có enrollment → trang chủ → lộ trình → bài học HybridView. Từ khu vực học có thư viện prompt và lịch sử. Instructor/admin → dashboard → lớp/chi tiết lớp, học viên, hoạt động và quản lý course.')
    table(d,['Luồng','Điều kiện vào','Kết quả'],[
        ['Đăng nhập','Supabase Auth có phiên Google hợp lệ','Đồng bộ public.users; xác định learner hoặc instructor'],
        ['Chọn lớp','Learner có enrollment hoặc lớp self_enroll','Gắn class đang học, nạp progress'],
        ['Mở bài','Class hợp lệ; lesson thuộc curriculum','HybridView hiển thị nội dung bài'],
        ['Chạy prompt','Prompt có nội dung, token và class hợp lệ ở chế độ live','Output, metadata model/latency; đánh giá và lưu attempt/progress'],
        ['Quản lý lớp','Instructor/admin','Quan sát tiến độ, sửa ghi danh/nội dung']
    ],[1.35,2.4,3.1])
    h(d,'2 Yêu cầu chức năng đang triển khai')
    table(d,['ID','Tác nhân và hành vi','Điều kiện/đầu ra quan sát được'],[
        ['FR01','Khách đăng nhập Google','LandingLoginScreen gọi Supabase OAuth; App resolve user'],
        ['FR02','Hệ thống phân vai','public.users.role learner → learner flow; instructor/admin → instructor shell'],
        ['FR03','Learner chọn lớp','Chỉ lớp được quyền hoặc công khai; self_enroll ghi danh rồi mở dashboard'],
        ['FR04','Learner mở lộ trình/bài','Lấy published curriculum từ DB; nếu lỗi có LABS_DATA fallback'],
        ['FR05','Learner đọc nội dung','Lý thuyết, tình huống, mục tiêu, gợi ý, lời giải; tutorial có trạng thái lưu DB'],
        ['FR06','Learner chạy prompt','Live gọi POST /api/generate; simulated chỉ khi chọn rõ trong cấu hình'],
        ['FR07','Learner nhận đánh giá','POST /api/evaluate cho 5 tiêu chí 0–2; tổng 0–10'],
        ['FR08','Learner xem lại','So sánh phiên bản; thư viện tìm/lọc/copy; lịch sử localStorage xuất Markdown'],
        ['FR09','Instructor xem dữ liệu','Dashboard, class detail, learner table, activity stream từ Supabase'],
        ['FR10','Instructor quản lý học viên','Thêm bằng email, sửa tên/trạng thái, remove khỏi class bằng trạng thái'],
        ['FR11','Instructor quản lý nội dung','Course publication; module, lesson, rubric, resource CRUD']
    ],[.55,2.15,4.15])
    h(d,'3 Luồng xử lý chi tiết')
    h(d,'3.1 Xác thực và truy cập lớp',2)
    nums(d,['Google OAuth tạo phiên Supabase; App đọc session và đồng bộ user profile.','Role trong public.users quyết định giao diện. Learner được resolve Learner → Enrollment còn hiệu lực → Class → Course/Client.','Nếu không có enrollment, App hiển thị chọn lớp. Class self_enroll có RPC/logic tự ghi danh; lớp thường cần instructor thêm trước.','API AI kiểm tra Bearer token, classId và quyền đọc class theo RLS. Lesson UUID phải thuộc course của class.'])
    h(d,'3.2 Chạy và đánh giá prompt',2)
    nums(d,['Learner chọn lesson, viết prompt trong HybridView và chọn chế độ live hoặc simulated.','Ở live, frontend gửi prompt, classId, lessonId, systemInstruction và inputContext đến /api/generate. Server ưu tiên OpenAI nếu có OPENAI_API_KEY; nếu không, dùng Gemini khi có GEMINI_API_KEY.','Frontend nhận output hoàn chỉnh, sau đó hiển thị dần bằng hiệu ứng chữ; đây không phải token stream trực tiếp từ model.','Frontend gửi prompt và output đến /api/evaluate để chấm 5 tiêu chí. Kết quả được kiểm tra theo contract.','Attempt và lesson_progress được ghi cùng giao dịch qua RPC record_prompt_attempt_and_progress khi dùng Supabase. Progress được nạp lại từ server.'])
    h(d,'3.3 Quản lý nội dung và học viên',2)
    nums(d,['Instructor mở dashboard/lớp; ứng dụng tổng hợp enrollment, attempt và progress.','Course có publication_status draft/published/archived; module và lesson cũng có trạng thái.','Instructor tạo/sửa/xóa module, lesson, rubric, resource trong CourseCurriculumEditor.','Thêm learner theo email tạo/đồng bộ user và enrollment. Remove chỉ đổi enrollment.status = removed.'])
    h(d,'4 Mô hình dữ liệu hiện tại')
    p(d,'Quan hệ chính: clients 1–N classes; courses 1–N classes; courses 1–N course_modules 1–N lessons; lessons 1–N rubric criteria/resources; users 1–1 learners; learners N–N classes qua enrollments; learner/class/lesson 1–N prompt_attempts và 1–1 lesson_progress.')
    table(d,['Bảng','Khóa và trường cốt lõi','Quan hệ / mục đích'],[
        ['users','id, email unique, role','Hồ sơ role learner/instructor/admin'],
        ['clients','id, name, industry','Tổ chức của class'],
        ['courses','id, status, publication_status, version','Nội dung dùng lại cho nhiều class'],
        ['classes','id, class_code unique, course_id, client_id, enrollment_mode','Một đợt triển khai course'],
        ['learners','id, learner_code unique, user_id unique','Hồ sơ học viên của user'],
        ['enrollments','learner_id + class_id unique, status','Quyền tham gia class'],
        ['course_modules / lessons','course_id / module_id, position, status','Cấu trúc curriculum và bài học'],
        ['lesson_rubric_criteria / lesson_resources','lesson_id, position','Rubric và tài nguyên từng bài'],
        ['course_instructors','course_id + user_id, permission','Gán instructor vào course'],
        ['prompt_attempts','learner_id, class_id, lesson_ref_id, prompt_text, ai_output, evaluation_json','Lịch sử lần chạy trên server'],
        ['lesson_progress','learner_id + class_id + lesson_id unique, status, attempts_count, best_score','Trạng thái học hiện tại']
    ],[1.75,2.8,2.3])
    h(d,'5 Công nghệ và tích hợp đã thấy trong mã')
    table(d,['Lớp','Công nghệ / trách nhiệm'],[
        ['Frontend','React 19, TypeScript, Vite 6, Tailwind CSS 3; App.tsx quản lý điều hướng theo state'],
        ['Xác thực và dữ liệu','Supabase Auth Google OAuth; Supabase PostgreSQL qua supabase-js; RLS và SQL migrations'],
        ['API','Hai handler /api/generate và /api/evaluate; kiểm tra Bearer token và quyền class/lesson'],
        ['AI','Server hỗ trợ OpenAI Responses API khi có OPENAI_API_KEY; nếu không dùng Gemini API khi có GEMINI_API_KEY'],
        ['Dữ liệu trình duyệt','sessionStorage cho điều hướng; localStorage cho cấu hình và màn lịch sử riêng'],
        ['Triển khai','Có URL Vercel do nhóm cung cấp; chưa xác minh commit/cấu hình thực tế']
    ],[1.55,5.3])
    h(d,'6 Phân quyền và quy tắc dữ liệu')
    bullets(d,['RLS của Supabase kiểm soát đọc/sửa theo role và enrollment; instructor có quyền quản lý nội dung theo policy.','API AI yêu cầu token và classId; bài lesson UUID phải thuộc course của class.','Course/module/lesson published quyết định nội dung có thể đọc; việc tồn tại file giáo trình không tự làm bài xuất hiện trong UI.','Dữ liệu progress server và lịch sử UI là hai đường dữ liệu khác nhau; không khẳng định lịch sử UI đồng bộ đa thiết bị.','Admin hiện dùng instructor shell; quyền admin riêng chưa được định nghĩa trong App.'])
    h(d,'7 Điều kiện nghiệm thu baseline và câu hỏi mở')
    table(d,['Kiểm tra baseline','Kết quả mong đợi / quyết định còn mở'],[
        ['Đăng nhập learner chưa có enrollment','Bị đưa tới chọn lớp, không khôi phục route bài học cũ'],
        ['Tự ghi danh class công khai','Chọn lớp thành công rồi mới mở dashboard'],
        ['Chạy live không có quyền class','API từ chối theo token/class RLS'],
        ['Chạy prompt và lưu progress','Attempt và progress cùng được ghi; reload lấy progress từ server'],
        ['Remove learner','Enrollment thành removed; user và learner còn tồn tại'],
        ['Commit deploy, dữ liệu course published, quyền admin riêng, lịch sử đa thiết bị','[Thầy/PM xác nhận trước nghiệm thu production]']
    ],[2.7,4.15])
    h(d,'8 Những tuyên bố không đưa vào baseline')
    p(d,'Mã main không chứng minh Google Sheets roster sync, Redis, worker đồng bộ 5 phút, schema email làm primary key, hoặc KPI cố định trong tài liệu cũ. Nếu dự án muốn các chức năng đó, cần tạo change request hoặc phần to-be riêng và được phê duyệt.')
    source(d,[('Ứng dụng và màn hình','src/App.tsx'),('Stack','package.json'),('Database lõi','supabase/schema.sql'),('Curriculum và bảng phụ','supabase/migrations/004_course_content_management.sql'),('Enrollment/RLS','supabase/migrations/006_enforce_course_enrollment_access.sql'),('Attempt/progress','supabase/migrations/016_add_server_backed_lesson_progress.sql'),('API authorization','src/services/apiAuthorizationService.ts'),('AI provider logic','src/services/apiServerService.ts'),('Frontend AI','src/services/llmService.ts')])
    out=OUT/'PRD_Promptify_Quan_ly_du_an.docx'; d.save(out); return out

if __name__=='__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    for path in (brd(),prd()): print(path)
