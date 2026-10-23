from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "document-drafts" / "Bao_cao_doi_chieu_Promptify_voi_PRD_2026-09-26.docx"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE_BLUE = "EEF4FA"
PALE_GRAY = "F5F7F9"
GRAY = "666666"
LIGHT_BORDER = "D9D9D9"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_col_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def style_cell_text(cell, size=9, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        for run in p.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.text = text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 120, 120, 120, 120)
        set_col_width(cell, widths[i])
        style_cell_text(cell, size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_col_width(cell, widths[i])
            if status_col is not None and i == status_col:
                lower = str(value).lower()
                if "chưa" in lower or "thiếu" in lower or "0%" in lower:
                    set_cell_shading(cell, RED)
                elif "một phần" in lower or "gần" in lower or "mô phỏng" in lower:
                    set_cell_shading(cell, AMBER)
                else:
                    set_cell_shading(cell, GREEN)
            elif row_idx % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or (status_col is not None and i == status_col) else WD_ALIGN_PARAGRAPH.LEFT
            style_cell_text(cell, size=8.5, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title_style.font.size = Pt(26)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(12)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size, before, after in (("Heading 1", 16, 14, 7), ("Heading 2", 12, 10, 5), ("Heading 3", 10.5, 8, 4)):
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = BLACK
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Bullet 2", "List Number"):
    st = styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("Báo cáo đối chiếu Promptify với PRD")

p = doc.add_paragraph()
p.add_run("Phạm vi đánh giá triển khai chức năng tiến độ và khoảng trống sản phẩm").bold = True
p.paragraph_format.space_after = Pt(18)

meta_rows = [
    ("Ngày đánh giá", "26 tháng 9 năm 2026"),
    ("PRD", "Prompt Engineering Learning Webapp phiên bản 1.0"),
    ("Mã nguồn", "Repository Promptify tại commit 06a4505 cùng các thay đổi chưa theo dõi tại thời điểm rà soát"),
    ("Mục đích", "Xác định phần đã làm, phần còn thiếu, phần làm ngoài PRD và mức độ hoàn thành so với đúng phạm vi PRD"),
]
add_table(doc, ["Thông tin", "Nội dung"], meta_rows, [1.55, 5.35])

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Kết luận chính")
r.bold = True
r.font.size = Pt(13)
p.paragraph_format.space_after = Pt(5)

p = doc.add_paragraph()
r = p.add_run("52 phần trăm")
r.bold = True
r.font.size = Pt(25)
r.font.color.rgb = RGBColor.from_string(NAVY)
p.add_run(" mức hoàn thành ước tính theo trọng số ưu tiên của PRD")
p.paragraph_format.space_after = Pt(8)

doc.add_paragraph(
    "Promptify đã hình thành một sản phẩm học Prompt Engineering có trải nghiệm learner và instructor khá sâu, đặc biệt ở playground, bài học, đánh giá AI, tiến độ và lịch sử thực hành. Tuy nhiên, sản phẩm chưa triển khai thật hai năng lực vận hành mà PRD đặt làm P0: đồng bộ roster từ Google Sheets và cưỡng chế hết hạn theo expiration timestamp. Vì vậy, tuy phạm vi MVP học tập có thể được xem là gần hoàn chỉnh theo tài liệu nội bộ, mức tuân thủ PRD gốc chỉ ở khoảng một nửa."
)

add_label_paragraph(doc, "Quyết định đề xuất ", "Không dùng nhãn Production Ready theo PRD trước khi hoàn thành roster sync, kiểm tra hết hạn phía server và loại bỏ các trạng thái đồng bộ mô phỏng khỏi giao diện.")

doc.add_page_break()

# Scope and method
doc.add_heading("Phạm vi và phương pháp đánh giá", level=1)
doc.add_paragraph(
    "Báo cáo đối chiếu chín yêu cầu chức năng trong mục Functional Requirements and Feature Prioritization của PRD với code, schema database, migration, test và tài liệu bàn giao đang có trong repository. Điểm số phản ánh mức triển khai có thể chứng minh từ mã nguồn, không phản ánh mức độ hấp dẫn của giao diện hoặc khối lượng công việc đã làm."
)

doc.add_heading("Cách tính tiến độ", level=2)
add_bullet(doc, "P0 Critical có trọng số 3 vì là điều kiện bắt buộc để phát hành theo PRD.")
add_bullet(doc, "P1 High có trọng số 2 vì ảnh hưởng trực tiếp đến vận hành và theo dõi lớp.")
add_bullet(doc, "P2 Medium có trọng số 1 vì có thể trì hoãn sau lần phát hành đầu.")
add_bullet(doc, "Mỗi yêu cầu được chấm từ 0 đến 100 dựa trên mức triển khai thật. Giao diện mô phỏng không được tính như tích hợp backend hoàn chỉnh.")
doc.add_paragraph("Công thức: tổng của điểm yêu cầu nhân trọng số, chia cho tổng điểm tối đa 2.200. Kết quả là 1.140 chia 2.200, tương đương 51,8 phần trăm và được làm tròn thành 52 phần trăm.")

spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(0)
doc.add_heading("Tóm tắt trạng thái", level=2)
summary_rows = [
    ("Gần hoàn thành", "2", "Interactive Prompt Playground; Prompt History and Export"),
    ("Hoàn thành một phần", "5", "Google OAuth; temporal access; Class ID mapping; expiration banner; telemetry"),
    ("Chưa triển khai thật", "2", "Google Sheets roster sync; manual sync status and force sync"),
]
add_table(doc, ["Nhóm", "Số yêu cầu", "Nội dung"], summary_rows, [1.55, 1.05, 4.30], status_col=0)

doc.add_heading("Bằng chứng xác minh", level=2)
doc.add_paragraph(
    "Lệnh npm run build đã hoàn thành thành công với TypeScript và Vite. Ba lệnh test được khai báo trong package.json không chạy hết trong phiên đánh giá vì tiến trình tsx gặp lỗi môi trường uv_os_get_passwd returned ENOMEM. Vì vậy, báo cáo dựa trên build thành công, kiểm tra tĩnh code và schema; chưa coi toàn bộ test suite là đã được xác nhận lại."
)

# Matrix
doc.add_heading("Ma trận đối chiếu yêu cầu PRD", level=1)
matrix_rows = [
    ("1", "P0", "Gmail SSO qua OAuth 2.0", "85", "Một phần", "Supabase Google OAuth, đồng bộ user và phân vai đã có. Người chưa ghi danh không vào được lớp. Chưa thấy ràng buộc Google Workspace domain riêng trong code."),
    ("2", "P0", "Google Sheet Roster Sync", "0", "Chưa làm", "Không có Google Sheets API, worker, lịch poll, Redis cache hoặc bảng trạng thái sync. Chỉ có câu chữ và nút mô phỏng trong banner."),
    ("3", "P0", "Temporal Expiration Enforcement", "25", "Một phần", "Có classes.end_date và enrollment status. RLS và API access chỉ kiểm tra status active, chưa so sánh thời gian hiện tại với end_date hoặc expiration timestamp trên mỗi học viên."),
    ("4", "P0", "Class ID Mapping", "85", "Một phần", "Class liên kết Course, Client, enrollment và curriculum. Người học chọn lớp đã được cấp quyền; không có bước nhập Class ID đúng như CUJ trong PRD."),
    ("5", "P0", "Interactive Prompt Playground", "95", "Gần hoàn thành", "Có system prompt, user prompt, temperature, live AI endpoint, output, token ước tính, latency, chấm điểm và nhiều workspace. Chưa thấy Top P được điều khiển trong luồng chính."),
    ("6", "P1", "Expiration Notice Banner", "20", "Mô phỏng", "Banner đúng màu PRD nhưng thời gian đếm từ expiryDurationHours cố định, không lấy end_date. Nội dung sync và last sync là state cục bộ giả lập."),
    ("7", "P1", "Manual Sync Status and Force Sync", "5", "Mô phỏng", "Có nút và trạng thái giao diện, nhưng handleForceSync chỉ setTimeout 800 ms; không gọi API hay ghi log sync."),
    ("8", "P1", "Telemetry and Session Logging", "65", "Một phần", "prompt_attempts lưu prompt, output, model, latency và evaluation; lesson_progress và dashboard instructor đã có. Chưa lưu token count và active session duration theo schema."),
    ("9", "P2", "Prompt History and Export", "90", "Gần hoàn thành", "Có lịch sử theo attempt, copy prompt và output, export Markdown. Một số view còn dùng localStorage nên độ bền dữ liệu phụ thuộc luồng sử dụng."),
]
add_table(doc, ["Số", "Mức", "Yêu cầu", "Điểm", "Trạng thái", "Nhận định và bằng chứng"], matrix_rows, [0.35, 0.45, 1.40, 0.48, 0.90, 3.32], status_col=4)

doc.add_paragraph(
    "Kết quả này không mâu thuẫn với tài liệu Promptify MVP Completion ngày 24 tháng 9 năm 2026. Tài liệu đó định nghĩa MVP theo learning loop, Supabase và AI evaluation; PRD gốc lại đặt Google Sheets sync và temporal expiration làm điều kiện P0. Hai tài liệu dùng hai định nghĩa hoàn thành khác nhau."
)

# What is done
doc.add_heading("Những phần Promptify đã làm được", level=1)
doc.add_heading("Nền tảng học và thực hành", level=2)
add_bullet(doc, "Playground hai cột cho phép nhập system instruction, user prompt, chỉnh temperature, chạy AI và xem output kèm latency và token ước tính.")
add_bullet(doc, "Luồng bài học có theory, bài tập, hint, lời giải, rubric, retry evaluation và so sánh nhiều lần thử.")
add_bullet(doc, "Curriculum đã mở rộng thành 22 module và có thêm các lab thực hành chuyên sâu cho PII, context engineering, structured output, sampling và prompt injection.")
add_bullet(doc, "Lịch sử thực hành cho phép xem lại, sao chép prompt, sao chép output và xuất Markdown.")

doc.add_heading("Xác thực và phân quyền cơ bản", level=2)
add_bullet(doc, "Đăng nhập Google được thực hiện qua Supabase Auth; hồ sơ người dùng được đồng bộ vào bảng users.")
add_bullet(doc, "Role learner, instructor và admin được lưu ở database; learner không tự nâng quyền.")
add_bullet(doc, "Enrollment, class, course và lesson được bảo vệ bằng RLS và kiểm tra phía API trước khi gọi mô hình AI.")
add_bullet(doc, "Lớp private yêu cầu ghi danh; lớp public testing hỗ trợ self enrollment qua RPC.")

doc.add_heading("Quản lý lớp và theo dõi học tập", level=2)
add_bullet(doc, "Instructor có dashboard, danh sách lớp, learner detail, activity stream, thêm hoặc xóa học viên và quản lý curriculum.")
add_bullet(doc, "prompt_attempts lưu prompt, output, evaluation, model, latency và thời điểm chạy; lesson_progress được đồng bộ bằng trigger và RPC.")
add_bullet(doc, "Progress được tách theo learner, class và lesson; ứng dụng có thể chuyển giữa nhiều lớp đã ghi danh.")

doc.add_heading("An toàn AI và dữ liệu", level=2)
add_bullet(doc, "API key được giữ phía server trong các endpoint generate và evaluate.")
add_bullet(doc, "Có PII detection và one click sanitization ở phía client trước khi gửi prompt.")
add_bullet(doc, "Có prompt structure analyzer, rubric năm tiêu chí và nội dung hướng dẫn phòng thủ prompt injection.")

# Missing
doc.add_heading("Những phần còn thiếu so với PRD", level=1)
missing_rows = [
    ("Google Sheets roster sync", "P0", "Không có tích hợp API v4, service account hoặc OAuth cho Sheet, worker 5 phút, mapping email class expiration, retry/backoff, Redis cache hay audit log.", "Người quản lý vẫn phải ghi danh bằng UI hoặc database; CUJ onboarding dưới 5 phút chưa được chứng minh."),
    ("Hết hạn theo timestamp", "P0", "Schema dùng end_date ở cấp class nhưng access helper chỉ kiểm tra status active. Không có job đổi enrollment sang expired và không có middleware so sánh now với end_date.", "Tài khoản có thể tiếp tục truy cập nếu class và enrollment vẫn active sau ngày kết thúc."),
    ("Class ID onboarding", "P0", "Luồng hiện tại chọn lớp đã được enroll; màn hình còn tuyên bố biết mã lớp không tự cấp quyền. Không có bước nhập Class ID để xác nhận như CUJ 2.", "Trải nghiệm khác PRD. Cần quyết định giữ mô hình bảo mật hiện tại hay thêm mã lớp như bước xác nhận không cấp quyền."),
    ("Expiration banner thật", "P1", "Countdown dùng số giờ cố định và chuỗi ngày hết hạn hard coded, không tính từ class end_date.", "Thông tin thời gian có thể sai và gây hiểu nhầm cho học viên."),
    ("Sync status và force sync", "P1", "Nút sync chỉ chạy setTimeout và thay chữ lastSync; không gọi backend.", "UI thể hiện một năng lực hệ thống chưa tồn tại."),
    ("Telemetry đầy đủ", "P1", "Có attempt và latency nhưng chưa có token_count trong bảng prompt_attempts, active session duration, login timestamp và execution cap theo người hoặc lớp.", "Chưa đo đủ các KPI engagement và chi phí trong PRD."),
    ("Kiểm soát chi phí và rate limit", "Security", "Server có xử lý lỗi provider 429 nhưng chưa có quota hoặc rate limit nội bộ theo student và Class ID.", "Rủi ro chi phí LLM vẫn còn dù API key được bảo vệ."),
    ("Redis cache và retry Sheets", "Kiến trúc", "Không có dependency hoặc code Redis; không có exponential backoff cho Sheets API.", "Technical Handover Specifications chưa được đáp ứng."),
]
add_table(doc, ["Khoảng trống", "Mức", "Thiếu gì", "Tác động"], missing_rows, [1.35, 0.65, 2.65, 2.55], status_col=None)

doc.add_heading("Điểm cần sửa ngay để tránh hiểu nhầm", level=2)
doc.add_paragraph(
    "ExpirationBanner hiện hiển thị Google Sheet Roster Đồng bộ và nút Đồng bộ lại. Hai thành phần này là mô phỏng hoàn toàn trong trình duyệt. Nếu sản phẩm được demo cho stakeholder, cần gắn nhãn demo rõ ràng hoặc ẩn chúng cho đến khi backend thật được kết nối. Đây là vấn đề trung thực trạng thái sản phẩm, không chỉ là một thiếu sót kỹ thuật."
)

# Extra
doc.add_heading("Những phần làm ngoài phạm vi PRD", level=1)
doc.add_paragraph(
    "Các hạng mục dưới đây không phải lỗi. Chúng tạo giá trị cho sản phẩm học tập, nhưng đã tiêu tốn phạm vi triển khai trong khi hai P0 vận hành vẫn chưa hoàn thành. Cần phân biệt phần vượt phạm vi hữu ích với phần nên hoãn để tập trung vào release criteria."
)
extra_rows = [
    ("Curriculum 22 module", "Hữu ích cao", "PRD chỉ yêu cầu bài tập theo lớp; repository có bộ giáo trình nhiều cấp độ và nhiều persona."),
    ("AI evaluation và rubric", "Hữu ích cao", "Có generate, evaluate, retry và năm tiêu chí chấm điểm, sâu hơn yêu cầu output feedback cơ bản."),
    ("Prompt structure analyzer", "Hữu ích cao", "Phân tích Role, Context, Task, Constraint, Format, Example và Evidence theo thời gian thực."),
    ("PII scrubber", "Hữu ích cao", "Phát hiện và thay thế CCCD, số tài khoản, số điện thoại và dữ liệu mẫu nhạy cảm."),
    ("A B compare và versioning", "Hữu ích cao", "So sánh prompt và output giữa các lần thử, giải thích cải thiện."),
    ("AI Coach và guided walkthrough", "Hữu ích", "Onboarding tương tác, hint và trợ lý theo context bài học."),
    ("Instructor content management", "Ngoài PRD", "Tạo sửa khóa học, module, lesson, rubric và resource; PRD chủ yếu mô tả monitoring."),
    ("Nhiều workspace", "Ngoài PRD", "Hybrid, Notebook, Playground và lesson workbench thay vì một playground duy nhất."),
    ("Public testing class", "Ngoài PRD", "Self enrollment cho lớp testing không được nêu trong CUJ gốc."),
    ("Prompt library và save prompt", "Mở rộng P2", "Không chỉ export lịch sử mà còn tái sử dụng prompt tốt trong thư viện."),
]
add_table(doc, ["Hạng mục", "Đánh giá", "Nhận định"], extra_rows, [1.75, 1.15, 4.30])

doc.add_paragraph(
    "Khuyến nghị giữ các phần ngoài PRD đã gắn trực tiếp với learning loop như evaluation, compare, PII scrubber và curriculum. Các mở rộng quản trị nội dung hoặc nhiều workspace nên tạm đóng băng cho đến khi hoàn thành roster sync, expiration enforcement và telemetry chi phí."
)

# Risks and roadmap
doc.add_heading("Rủi ro phát hành theo PRD", level=1)
risk_rows = [
    ("Critical", "Truy cập sau ngày hết hạn", "RLS chỉ kiểm tra status active, không kiểm tra thời gian kết thúc", "Thêm kiểm tra now giữa start_date và end_date trong helper RLS và API; có job cập nhật expired."),
    ("Critical", "Roster không tự đồng bộ", "Không có Google Sheets integration", "Xây sync service, mapping chuẩn, idempotency, retry và audit log."),
    ("High", "UI báo sync thành công giả", "Force sync chỉ setTimeout", "Ẩn hoặc gắn nhãn demo trước khi có backend."),
    ("High", "Không kiểm soát chi phí theo lớp", "Chưa có token cap hoặc rate limit nội bộ", "Thêm quota theo learner và class, trả 429 có thông báo rõ."),
    ("Medium", "KPI không đủ dữ liệu", "Thiếu token count và session duration", "Bổ sung event hoặc cột telemetry, thống nhất nguồn dữ liệu server."),
]
add_table(doc, ["Mức", "Rủi ro", "Nguyên nhân", "Xử lý đề xuất"], risk_rows, [0.70, 1.55, 2.30, 2.65], status_col=0)

doc.add_heading("Lộ trình đề xuất để đạt PRD", level=1)
roadmap_rows = [
    ("Giai đoạn 1", "Độ tin cậy trạng thái", "Ẩn mock sync; banner dùng end_date thật; thêm test thời gian; sửa wording Production Ready.", "Có thể demo trung thực"),
    ("Giai đoạn 2", "P0 access control", "Cập nhật RLS và API để kiểm tra start_date, end_date và enrollment status trên mọi request.", "Chặn truy cập hết hạn"),
    ("Giai đoạn 3", "Google Sheets sync", "Tích hợp Sheets API, worker định kỳ, mapping idempotent, retry backoff, audit log và nút force sync thật.", "Hoàn thành CUJ quản trị roster"),
    ("Giai đoạn 4", "Telemetry và cost control", "Lưu token count, session duration, last login; thêm quota và rate limit theo learner và class.", "Đo KPI và kiểm soát chi phí"),
    ("Giai đoạn 5", "Nghiệm thu", "Chạy test integration với Sheet thật, clock boundary, expired JWT, RLS, force sync và dashboard telemetry.", "Sẵn sàng đánh giá lại"),
]
add_table(doc, ["Giai đoạn", "Mục tiêu", "Phạm vi", "Kết quả"], roadmap_rows, [0.85, 1.35, 3.55, 1.45])

doc.add_heading("Tiêu chí để nâng tiến độ lên 100 phần trăm", level=2)
add_number(doc, "Google Sheet mới hoặc thay đổi phải tạo, cập nhật và vô hiệu hóa enrollment đúng trong thời gian cam kết, có log và retry.")
add_number(doc, "Mọi request đọc curriculum, chạy AI và ghi attempt phải bị chặn khi thời gian lớp hoặc enrollment đã hết hạn.")
add_number(doc, "Banner phải lấy timestamp thật, hiển thị đúng múi giờ và không có trạng thái sync giả.")
add_number(doc, "Instructor phải thấy last sync thật, lỗi sync và nút force sync có kết quả từ backend.")
add_number(doc, "Telemetry phải lưu execution count, token usage, latency và active session duration theo learner và class.")
add_number(doc, "Rate limit và token cap phải được kiểm tra phía server trước khi gọi provider.")
add_number(doc, "Test tự động và test integration phải chạy thành công trong môi trường staging có Supabase và Google Sheets cấu hình thật.")

# Evidence appendix
doc.add_page_break()
doc.add_heading("Phụ lục bằng chứng kỹ thuật", level=1)
evidence_rows = [
    ("Google OAuth và session", "src/App.tsx; src/services/dbService.ts; src/services/supabaseClient.ts"),
    ("API authorization", "src/services/apiAuthorizationService.ts; api/generate.ts; api/evaluate.ts"),
    ("Enrollment và RLS", "supabase/schema.sql; supabase/migrations/006_enforce_course_enrollment_access.sql"),
    ("Expiration và sync UI", "src/components/common/ExpirationBanner.tsx"),
    ("Playground", "src/components/playground/PlaygroundView.tsx; src/components/hybrid/HybridView.tsx; src/services/llmService.ts"),
    ("Prompt history và export", "src/components/history/PromptHistoryView.tsx; src/components/telemetry/HistoryDrawer.tsx"),
    ("Attempt và progress", "supabase/migrations/003_create_prompt_attempts.sql; supabase/migrations/016_add_server_backed_lesson_progress.sql; src/services/dbService.ts"),
    ("Instructor monitoring", "src/components/instructor/InstructorDashboard.tsx; src/components/instructor/ActivityStreamView.tsx"),
    ("Phần vượt PRD", "curriculum-levels; curriculum-specs; src/services/promptStructureAnalyzer.ts; src/services/labComplianceService.ts"),
    ("Build và test scripts", "package.json; TEST_PLAN_AND_EXECUTION_LOG.md"),
]
add_table(doc, ["Chủ đề", "File đối chiếu chính"], evidence_rows, [2.15, 5.05])

doc.add_heading("Giới hạn của đánh giá", level=2)
doc.add_paragraph(
    "Báo cáo đánh giá những gì có thể chứng minh trong repository tại thời điểm rà soát. Không có quyền truy cập vào cấu hình Supabase production, Google Cloud project, biến môi trường Vercel hoặc log vận hành thật. Nếu có worker hoặc dịch vụ sync nằm ngoài repository, điểm Google Sheets và expiration cần được cập nhật sau khi cung cấp code, cấu hình và bằng chứng chạy staging."
)

# Footer
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Promptify PRD Gap Assessment  |  26 09 2026")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)

doc.core_properties.title = "Báo cáo đối chiếu Promptify với PRD"
doc.core_properties.subject = "Đánh giá phần đã làm phần còn thiếu phần vượt phạm vi và tiến độ PRD"
doc.core_properties.author = "Promptify Project Review"
doc.core_properties.keywords = "Promptify, PRD, gap analysis, progress"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
